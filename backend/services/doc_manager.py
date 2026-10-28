import os
import time
import json
import uuid
import hashlib
import asyncio
import logging
import re
from typing import Dict, Any, List, Optional, Tuple
import pandas as pd
import openpyxl

from qdrant_client import models, QdrantClient
from langchain_core.documents import Document
from langchain_text_splitters import MarkdownHeaderTextSplitter, RecursiveCharacterTextSplitter

from backend.core.config import (
    DOCS_DIR, KNOWLEDGE_DIR, FAQ_DIR, TELECOM_DIR, NAVIGATION_DIR, STATE_FILE,
    COLLECTION_NAME, CATEGORIES, NAVIGATION_MAP_FILE, APP_GUIDE_FILE, FAQ_EXCEL_FILE,
    TELECOM_EXCEL_FILE, TELECOM_PACKAGES_FILE
)
from backend.core.state import app_state

logger = logging.getLogger(__name__)

# Lock đồng bộ hóa các thao tác index tránh xung đột tài nguyên
indexing_lock = asyncio.Lock()


def get_file_hash(filepath: str) -> str:
    """Tính mã MD5 hash của file."""
    hasher = hashlib.md5()
    with open(filepath, "rb") as f:
        while chunk := f.read(65536):
            hasher.update(chunk)
    return hasher.hexdigest()


def load_state() -> dict:
    """Đọc file state lưu thông tin hash và id các chunk đã index."""
    if os.path.exists(STATE_FILE):
        try:
            with open(STATE_FILE, "r", encoding="utf-8") as f:
                return json.loads(f.read())
        except Exception as e:
            logger.error("[DocManager] Lỗi đọc state: %s", e)
    return {}


def save_state(state: dict):
    """Lưu state vào file."""
    content = json.dumps(state, ensure_ascii=False, indent=2)
    with open(STATE_FILE, "w", encoding="utf-8") as f:
        f.write(content)


def convert_telecom_excel_to_json(
    excel_path: str = TELECOM_EXCEL_FILE,
    json_path: str = TELECOM_PACKAGES_FILE
) -> int:
    """Chuyển đổi dữ liệu gói cước từ Excel sang JSON (goi_cuoc.xlsx -> goi_cuoc.json)."""
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Không tìm thấy file Excel tại '{excel_path}'")

    df = pd.read_excel(excel_path).fillna("")
    json_data = []
    for index, row in df.iterrows():
        metadata = {}
        for col in df.columns:
            val = row[col]
            if col == "gia_vnd" and val != "":
                try:
                    metadata[col] = int(float(val))
                except (ValueError, TypeError):
                    metadata[col] = str(val)
            elif isinstance(val, float) and val.is_integer():
                metadata[col] = str(int(val))
            else:
                metadata[col] = str(val) if val != "" else ""

        json_data.append({
            "id": f"pkg_{index + 1:03d}",
            "metadata": metadata
        })

    os.makedirs(os.path.dirname(os.path.abspath(json_path)), exist_ok=True)
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(json_data, f, ensure_ascii=False, indent=2)
    return len(json_data)


def refresh_domain_caches(rel_path: str = ""):
    """Đồng bộ lại cache gói cước và sơ đồ điều hướng khi file tương ứng thay đổi."""
    lower = rel_path.lower()
    if not lower or "goi_cuoc" in lower:
        if (not lower or lower.endswith((".xlsx", ".xls"))) and os.path.exists(TELECOM_EXCEL_FILE):
            try:
                convert_telecom_excel_to_json(TELECOM_EXCEL_FILE, TELECOM_PACKAGES_FILE)
            except Exception as e:
                logger.warning("[DocManager] Lỗi đồng bộ Excel sang JSON cho gói cước: %s", e)
        from backend.tools.telecom_packages import load_package_data
        app_state.telecom_packages = load_package_data()

    if not lower or "navigation_map" in lower:
        from backend.tools.app_navigation import load_navigation_map
        app_state.nav_map = load_navigation_map()


def determine_category(relative_path: str) -> str:
    """Xác định category dựa trên đường dẫn thư mục con."""
    norm = relative_path.replace("\\", "/").lower()
    if norm.startswith("knowledge/"):
        return "knowledge"
    elif norm.startswith("faq/"):
        return "faq"
    elif norm.startswith("telecom/"):
        return "telecom"
    elif norm.startswith("navigation/"):
        return "navigation"

    # Fallback theo tên file nếu ở thư mục gốc
    fname = os.path.basename(norm)
    if fname.startswith("goi_cuoc"):
        return "telecom"
    if fname in (os.path.basename(FAQ_EXCEL_FILE), "faq.xlsx"):
        return "faq"
    if fname in (os.path.basename(NAVIGATION_MAP_FILE), os.path.basename(APP_GUIDE_FILE)):
        return "navigation"
    return "knowledge"


def is_indexable_doc(relative_path: str) -> bool:
    """
    Kiểm tra file có hợp lệ để index vào vector database Qdrant hay không.
    - knowledge/: Tài liệu Markdown (.md)
    - faq/: Bảng hỏi đáp Excel (.xlsx, .xls)
    - navigation/: Hướng dẫn tác vụ Markdown (.md)
    Loại trừ thư mục telecom/, file json và file state.
    """
    norm = relative_path.replace("\\", "/").lower()
    fname = os.path.basename(norm)

    if fname == os.path.basename(STATE_FILE):
        return False
    if norm.startswith("telecom/") or fname.startswith("goi_cuoc."):
        return False
    if norm.endswith(".json"):
        return False

    if norm.startswith("faq/") or fname.endswith((".xlsx", ".xls")):
        return norm.endswith((".xlsx", ".xls"))

    return norm.endswith(".md")


def find_qa_columns(df: pd.DataFrame, allow_fallback: bool = True) -> Tuple[Optional[str], Optional[str], Optional[str]]:
    """
    Tự động nhận diện cột Câu hỏi, Câu trả lời và Từ khóa trong file Excel.
    Hỗ trợ cả tiếng Việt có dấu, không dấu và tiếng Anh.
    """
    cols = list(df.columns)
    q_col, a_col, kw_col = None, None, None
    for c in cols:
        c_clean = str(c).lower().strip().replace("_", " ").replace("-", " ")
        if (c_clean == "q" or any(term in c_clean for term in ["question", "câu hỏi", "cau hoi", "hỏi", "hoi"])) and not q_col:
            q_col = c
        elif (c_clean == "a" or any(term in c_clean for term in ["answer", "câu trả lời", "cau tra loi", "trả lời", "tra loi", "đáp", "dap"])) and not a_col:
            a_col = c
        elif any(term in c_clean for term in ["keyword", "từ khóa", "tu khoa", "tags"]) and not kw_col:
            kw_col = c

    # Dự phòng: Nếu có từ 2 cột trở lên và cho phép fallback, lấy cột đầu là Câu hỏi, cột hai là Câu trả lời
    if allow_fallback and (not q_col or not a_col) and len(cols) >= 2:
        q_col = cols[0]
        a_col = cols[1]

    return q_col, a_col, kw_col


def init_splitters():
    """Khởi tạo các bộ chia tài liệu Markdown."""
    headers_to_split_on = [
        ("#", "H1"),
        ("##", "H2"),
        ("###", "H3"),
    ]
    markdown_splitter = MarkdownHeaderTextSplitter(headers_to_split_on=headers_to_split_on, strip_headers=False)
    parent_splitter = RecursiveCharacterTextSplitter(chunk_size=3000, chunk_overlap=200)
    child_splitter = RecursiveCharacterTextSplitter(chunk_size=300, chunk_overlap=30)
    return markdown_splitter, parent_splitter, child_splitter


def split_markdown_documents_with_tables(docs, max_rows=10):
    """Tách các bảng Markdown dài giữ nguyên header."""
    new_docs = []
    for doc in docs:
        lines = doc.page_content.split("\n")
        i = 0
        current_text = []
        while i < len(lines):
            line = lines[i].strip()
            if line.startswith("|") and line.endswith("|") and i + 1 < len(lines):
                next_line = lines[i + 1].strip()
                if next_line.startswith("|") and next_line.endswith("|") and re.match(r'^\|[\s\-\|:]+\|$', next_line):
                    if current_text:
                        text = "\n".join(current_text).strip()
                        if text:
                            new_docs.append(Document(page_content=text, metadata=doc.metadata))
                        current_text = []

                    header = lines[i]
                    separator = lines[i + 1]
                    i += 2
                    rows = []
                    while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                        rows.append(lines[i])
                        i += 1

                    for j in range(0, len(rows), max_rows):
                        chunk_rows = rows[j:j + max_rows]
                        table_chunk = [header, separator] + chunk_rows
                        new_docs.append(Document(page_content="\n".join(table_chunk), metadata=doc.metadata))
                    continue
            current_text.append(lines[i])
            i += 1

        if current_text:
            text = "\n".join(current_text).strip()
            if text:
                new_docs.append(Document(page_content=text, metadata=doc.metadata))
    return new_docs


def index_single_markdown_file(
    filepath: str,
    stored_path_key: str,
    client: QdrantClient,
    store,
    dense_model,
    sparse_model,
    markdown_splitter,
    parent_splitter,
    child_splitter
) -> Tuple[List[str], List[str]]:
    """Xử lý và index một file Markdown."""
    filename = os.path.basename(filepath)
    with open(filepath, "r", encoding="utf-8") as f:
        content = f.read()

    normalized_lines = []
    for line in content.split("\n"):
        stripped = line.lstrip()
        if stripped.startswith(("# ", "## ", "### ")):
            normalized_lines.append(stripped)
        else:
            normalized_lines.append(line)
    content = "\n".join(normalized_lines)

    md_docs = markdown_splitter.split_text(content)
    md_docs = split_markdown_documents_with_tables(md_docs, max_rows=10)
    parent_chunks = parent_splitter.split_documents(md_docs)

    parent_docs = []
    for doc in parent_chunks:
        headers = []
        for h in ["H1", "H2", "H3"]:
            if h in doc.metadata:
                headers.append(doc.metadata[h])
        heading_context = " > ".join(headers)
        raw_text = f"[{heading_context}]\n\n{doc.page_content}" if heading_context else doc.page_content
        meta = {"file_path": stored_path_key, "filename": filename}
        screen_m = re.search(r'required_screen_id`?\)?:\s*`?([A-Z0-9_]+)`?', doc.page_content)
        if screen_m:
            meta["required_screen_id"] = screen_m.group(1)

        parent_docs.append(Document(
            page_content=raw_text,
            metadata=meta
        ))

    doc_ids = [str(uuid.uuid4()) for _ in parent_docs]

    child_items = []
    for i, parent_doc in enumerate(parent_docs):
        _id = doc_ids[i]
        _sub_docs = child_splitter.split_documents([parent_doc])
        for _doc in _sub_docs:
            child_items.append((_id, _doc.page_content))

    if not child_items:
        if doc_ids:
            store.mset(list(zip(doc_ids, parent_docs)))
        return doc_ids, []

    all_texts = [item[1] for item in child_items]

    # Dense embedding
    dense_vectors = []
    batch_size = 32
    for b_idx in range(0, len(all_texts), batch_size):
        batch = all_texts[b_idx:b_idx + batch_size]
        for attempt in range(3):
            try:
                vectors = dense_model.embed_documents(batch)
                dense_vectors.extend(vectors)
                break
            except Exception as e:
                if attempt == 2:
                    raise e
                time.sleep(2 ** attempt)

    # Sparse embedding
    sparse_results = list(sparse_model.embed(all_texts))

    points = []
    point_ids = []
    for (pid, text), dense_vec, sparse_res in zip(child_items, dense_vectors, sparse_results):
        child_id = str(uuid.uuid4())
        points.append(
            models.PointStruct(
                id=child_id,
                payload={"page_content": text, "metadata": {"doc_id": pid, "file_path": stored_path_key, "filename": filename}},
                vector={
                    "dense": dense_vec,
                    "sparse": models.SparseVector(
                        indices=sparse_res.indices.tolist(),
                        values=sparse_res.values.tolist()
                    )
                }
            )
        )
        point_ids.append(child_id)

    if points:
        client.upload_points(collection_name=COLLECTION_NAME, points=points)
    if doc_ids:
        store.mset(list(zip(doc_ids, parent_docs)))

    return doc_ids, point_ids


def index_single_excel_file(
    filepath: str,
    stored_path_key: str,
    client: QdrantClient,
    store,
    dense_model,
    sparse_model
) -> Tuple[List[str], List[str]]:
    """
    Xử lý và index một file Excel Q&A:
    - Row-level chunking: Mỗi dòng (1 cặp Q&A) là 1 chunk độc lập.
    - Vector Embedding: Chỉ tạo cho cột Câu hỏi (hoặc Câu hỏi + Từ khóa). Cột Câu trả lời lưu vào metadata.
    - Khi truy vấn, Qdrant so khớp câu hỏi của user với Câu hỏi mẫu, lấy nguyên văn Câu trả lời đưa vào context cho LLM.
    """
    filename = os.path.basename(filepath)
    df = pd.read_excel(filepath).fillna("")

    q_col, a_col, kw_col = find_qa_columns(df)
    if not q_col or not a_col:
        logger.warning("[DocManager] File %s không tìm thấy các cột Q&A hợp lệ. Bỏ qua.", filepath)
        return [], []

    parent_docs = []
    child_items = []

    for idx, row in df.iterrows():
        q_text = str(row[q_col]).strip()
        a_text = str(row[a_col]).strip()
        kw_text = str(row[kw_col]).strip() if kw_col else ""

        if not q_text and not a_text:
            continue

        parent_id = str(uuid.uuid4())
        embed_text = f"{q_text} ({kw_text})" if kw_text else q_text

        # Nội dung parent chunk chứa nguyên văn câu trả lời cho ngữ cảnh của LLM
        raw_text = f"[Hỏi & Đáp - {filename} - Dòng {idx + 1}]\nCâu hỏi: {q_text}\nCâu trả lời: {a_text}"

        parent_doc = Document(
            page_content=raw_text,
            metadata={
                "file_path": stored_path_key,
                "filename": filename,
                "row_idx": int(idx + 1),
                "question": q_text,
                "answer": a_text,
                "type": "qa_excel"
            }
        )
        parent_docs.append(parent_doc)
        child_items.append((parent_id, embed_text, q_text, a_text, int(idx + 1)))

    if not child_items:
        return [], []

    all_texts = [item[1] for item in child_items]

    # Dense embedding (Chỉ nhúng cột câu hỏi)
    dense_vectors = []
    batch_size = 32
    for b_idx in range(0, len(all_texts), batch_size):
        batch = all_texts[b_idx:b_idx + batch_size]
        for attempt in range(3):
            try:
                vectors = dense_model.embed_documents(batch)
                dense_vectors.extend(vectors)
                break
            except Exception as e:
                if attempt == 2:
                    raise e
                time.sleep(2 ** attempt)

    # Sparse embedding (Chỉ nhúng cột câu hỏi)
    sparse_results = list(sparse_model.embed(all_texts))

    points = []
    point_ids = []
    doc_ids = [item[0] for item in child_items]

    for (pid, embed_text, q_text, a_text, row_idx), dense_vec, sparse_res in zip(child_items, dense_vectors, sparse_results):
        child_id = str(uuid.uuid4())
        points.append(
            models.PointStruct(
                id=child_id,
                payload={
                    "page_content": embed_text,
                    "metadata": {
                        "doc_id": pid,
                        "file_path": stored_path_key,
                        "filename": filename,
                        "row_idx": row_idx,
                        "question": q_text,
                        "answer": a_text,
                        "type": "qa_excel"
                    }
                },
                vector={
                    "dense": dense_vec,
                    "sparse": models.SparseVector(
                        indices=sparse_res.indices.tolist(),
                        values=sparse_res.values.tolist()
                    )
                }
            )
        )
        point_ids.append(child_id)

    if points:
        client.upload_points(collection_name=COLLECTION_NAME, points=points)
    if doc_ids:
        store.mset(list(zip(doc_ids, parent_docs)))

    return doc_ids, point_ids


def delete_file_from_index(rel_path: str, file_info: dict, client: QdrantClient, store):
    """Xóa dữ liệu một file khỏi database và store."""
    doc_ids = file_info.get("doc_ids", [])
    point_ids = file_info.get("point_ids", [])

    if doc_ids and store:
        try:
            store.mdelete(doc_ids)
        except Exception as e:
            logger.warning("[DocManager] Lỗi xóa parent docs %s: %s", rel_path, e)

    if point_ids and client:
        try:
            client.delete(
                collection_name=COLLECTION_NAME,
                points_selector=models.PointIdsList(points=point_ids)
            )
        except Exception as e:
            logger.warning("[DocManager] Lỗi xóa points theo ID %s: %s", rel_path, e)

    if client:
        try:
            client.delete(
                collection_name=COLLECTION_NAME,
                points_selector=models.FilterSelector(
                    filter=models.Filter(
                        must=[
                            models.FieldCondition(
                                key="metadata.file_path",
                                match=models.MatchValue(value=rel_path)
                            )
                        ]
                    )
                )
            )
        except Exception:
            pass


class DocumentManager:
    """Service quản lý và cập nhật tài liệu đa thư mục trong docs/."""

    def format_size(self, size_bytes: int) -> str:
        """Đổi kích thước bytes sang chuỗi hiển thị dễ đọc."""
        if size_bytes < 1024:
            return f"{size_bytes} B"
        elif size_bytes < 1024 * 1024:
            return f"{size_bytes / 1024:.1f} KB"
        else:
            return f"{size_bytes / (1024 * 1024):.1f} MB"

    def scan_all_files(self) -> List[Dict[str, Any]]:
        """Quét toàn bộ file trong docs/ bao gồm các thư mục con."""
        results = []
        if not os.path.exists(DOCS_DIR):
            return results

        for root, dirs, files in os.walk(DOCS_DIR):
            for fname in sorted(files):
                if fname == os.path.basename(STATE_FILE):
                    continue
                fpath = os.path.join(root, fname)
                rel_path = os.path.relpath(fpath, DOCS_DIR).replace("\\", "/")
                results.append({
                    "full_path": fpath,
                    "relative_path": rel_path,
                    "filename": fname,
                    "category": determine_category(rel_path)
                })
        return results

    def resolve_file_path(self, identifier: str) -> Optional[str]:
        """Tìm đường dẫn tuyệt đối của file theo relative_path hoặc filename với cơ chế chống Path Traversal."""
        cleaned = identifier.replace("\\", "/").strip("/")
        abs_docs = os.path.abspath(DOCS_DIR)

        # 1. Thử trực tiếp dưới DOCS_DIR
        direct = os.path.abspath(os.path.join(DOCS_DIR, cleaned))
        if os.path.commonpath([abs_docs, direct]) == abs_docs and os.path.exists(direct) and os.path.isfile(direct):
            return direct

        # 2. Tìm trong các subfolder
        fname = os.path.basename(cleaned)
        for sub in CATEGORIES:
            cand = os.path.abspath(os.path.join(DOCS_DIR, sub, fname))
            if os.path.commonpath([abs_docs, cand]) == abs_docs and os.path.exists(cand) and os.path.isfile(cand):
                return cand
        return None

    def list_documents(self) -> List[Dict[str, Any]]:
        """Lấy danh sách toàn bộ tài liệu phân nhóm theo category kèm trạng thái index."""
        state = load_state()
        all_files = self.scan_all_files()
        docs = []

        for item in all_files:
            fpath = item["full_path"]
            rel_path = item["relative_path"]
            fname = item["filename"]
            category = item["category"]

            stat = os.stat(fpath)
            size = stat.st_size
            mtime = stat.st_mtime

            lower = fname.lower()
            if lower.endswith(".md"):
                ftype = "markdown"
            elif lower.endswith((".xlsx", ".xls")):
                ftype = "excel"
            elif lower.endswith(".json"):
                ftype = "json"
            else:
                ftype = "other"

            is_idxable = is_indexable_doc(rel_path)

            # Kiểm tra trạng thái index theo relative_path
            state_info = state.get(rel_path)
            is_indexed = bool(state_info and is_idxable)
            chunk_count = len(state_info.get("doc_ids", [])) if state_info else 0

            row_count = None
            if ftype == "excel":
                # Đọc số dòng nhanh với openpyxl read_only thay vì parse toàn bộ dataframe với pandas
                try:
                    wb = openpyxl.load_workbook(fpath, read_only=True)
                    sheet = wb.active
                    row_count = max(0, (sheet.max_row or 1) - 1)
                    wb.close()
                except Exception:
                    row_count = None

            docs.append({
                "filename": fname,
                "relative_path": rel_path,
                "category": category,
                "file_type": ftype,
                "size_bytes": size,
                "size_formatted": self.format_size(size),
                "modified_at": mtime,
                "modified_at_str": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(mtime)),
                "is_indexable": is_idxable,
                "is_indexed": is_indexed,
                "chunk_count": chunk_count,
                "row_count": row_count
            })

        return docs

    def get_document_content(self, identifier: str) -> Dict[str, Any]:
        """Đọc nội dung chi tiết của một tài liệu."""
        filepath = self.resolve_file_path(identifier)
        if not filepath:
            raise FileNotFoundError(f"Không tìm thấy file: {identifier}")

        rel_path = os.path.relpath(filepath, DOCS_DIR).replace("\\", "/")
        fname = os.path.basename(filepath)
        category = determine_category(rel_path)
        lower = fname.lower()

        if lower.endswith((".xlsx", ".xls")):
            df = pd.read_excel(filepath).fillna("")
            cols = [str(c) for c in df.columns]
            rows = []
            for _, r in df.iterrows():
                rows.append({str(c): str(r[c]) for c in cols})
            q_col, a_col, kw_col = find_qa_columns(df, allow_fallback=(category == "faq"))
            return {
                "type": "excel",
                "filename": fname,
                "relative_path": rel_path,
                "category": category,
                "columns": cols,
                "rows": rows,
                "total_rows": len(rows),
                "is_qa": bool(category == "faq" and q_col and a_col),
                "q_col": q_col,
                "a_col": a_col,
                "kw_col": kw_col
            }
        else:
            with open(filepath, "r", encoding="utf-8") as f:
                text = f.read()
            return {
                "type": "text",
                "filename": fname,
                "relative_path": rel_path,
                "category": category,
                "content": text
            }

    async def save_document_content(self, identifier: str, data: Dict[str, Any]) -> Dict[str, Any]:
        """Lưu nội dung chỉnh sửa của file và tự động re-index."""
        filepath = self.resolve_file_path(identifier)
        if not filepath:
            raise FileNotFoundError(f"Không tìm thấy file: {identifier}")

        fname = os.path.basename(filepath)
        rel_path = os.path.relpath(filepath, DOCS_DIR).replace("\\", "/")
        lower = fname.lower()

        def _write_file():
            if lower.endswith((".xlsx", ".xls")):
                rows = data.get("rows", [])
                cols = data.get("columns", [])
                if not cols and rows:
                    cols = list(rows[0].keys())
                df = pd.DataFrame(rows, columns=cols if cols else None)
                df.to_excel(filepath, index=False)
            else:
                content = data.get("content", "")
                with open(filepath, "w", encoding="utf-8") as f:
                    f.write(content)

        await asyncio.to_thread(_write_file)

        # Index lại
        index_result = await self.index_single_document(rel_path)

        await asyncio.to_thread(refresh_domain_caches, rel_path)

        return {
            "message": f"Đã lưu và cập nhật index cho {rel_path}",
            "relative_path": rel_path,
            "filename": fname,
            "index_result": index_result
        }

    async def upload_file(self, filename: str, content_bytes: bytes, category: Optional[str] = None) -> Dict[str, Any]:
        """Lưu file tải lên vào đúng subfolder category và tự động index an toàn."""
        clean_name = os.path.basename(filename.replace("\\", "/"))

        # Xác định thư mục đích
        if category in CATEGORIES:
            target_sub = category
        else:
            target_sub = determine_category(clean_name)

        target_dir = os.path.abspath(os.path.join(DOCS_DIR, target_sub))
        os.makedirs(target_dir, exist_ok=True)
        dest_filepath = os.path.join(target_dir, clean_name)

        def _write():
            with open(dest_filepath, "wb") as f:
                f.write(content_bytes)

        await asyncio.to_thread(_write)

        rel_path = f"{target_sub}/{clean_name}"
        index_result = await self.index_single_document(rel_path)

        await asyncio.to_thread(refresh_domain_caches, rel_path)

        return {
            "message": f"Đã tải lên vào '{target_sub}/' và cập nhật index",
            "filename": clean_name,
            "relative_path": rel_path,
            "category": target_sub,
            "index_result": index_result
        }

    async def delete_document(self, identifier: str) -> Dict[str, Any]:
        """Xóa file khỏi thư mục và un-index khỏi database."""
        filepath = self.resolve_file_path(identifier)
        if not filepath:
            raise FileNotFoundError(f"Không tìm thấy file: {identifier}")

        rel_path = os.path.relpath(filepath, DOCS_DIR).replace("\\", "/")
        fname = os.path.basename(filepath)

        async with indexing_lock:
            state = await asyncio.to_thread(load_state)
            client = app_state.qdrant_client
            store = app_state.store

            if rel_path in state:
                if client and store:
                    await asyncio.to_thread(delete_file_from_index, rel_path, state[rel_path], client, store)
                del state[rel_path]

            await asyncio.to_thread(save_state, state)

            if os.path.exists(filepath):
                await asyncio.to_thread(os.remove, filepath)

        await asyncio.to_thread(refresh_domain_caches, rel_path)

        return {"message": f"Đã xóa thành công {rel_path}", "relative_path": rel_path, "filename": fname}

    async def index_single_document(self, identifier: str) -> Dict[str, Any]:
        """Index một tài liệu cụ thể bằng in-process models và client."""
        filepath = self.resolve_file_path(identifier)
        if not filepath:
            return {"status": "error", "reason": f"Không tìm thấy file: {identifier}"}

        rel_path = os.path.relpath(filepath, DOCS_DIR).replace("\\", "/")
        fname = os.path.basename(filepath)

        if not is_indexable_doc(rel_path):
            return {"status": "skipped", "reason": f"File {rel_path} không thuộc diện vector index (dữ liệu gói cước hoặc json)."}

        # Đợi các model nạp xong nếu đang nạp
        if app_state.load_task and not app_state.load_task.done():
            await app_state.load_task

        client = app_state.qdrant_client
        store = app_state.store
        dense_model = app_state.dense_model
        sparse_model = app_state.sparse_model

        if not client or not store or not dense_model or not sparse_model:
            return {"status": "error", "reason": "Các thành phần AI chưa sẵn sàng để index."}

        async with indexing_lock:
            state = await asyncio.to_thread(load_state)
            # Xóa points cũ nếu có
            if rel_path in state:
                await asyncio.to_thread(delete_file_from_index, rel_path, state[rel_path], client, store)
                del state[rel_path]

            lower = fname.lower()
            if lower.endswith((".xlsx", ".xls")):
                doc_ids, point_ids = await asyncio.to_thread(
                    index_single_excel_file,
                    filepath, rel_path, client, store, dense_model, sparse_model
                )
            else:
                markdown_splitter, parent_splitter, child_splitter = init_splitters()
                doc_ids, point_ids = await asyncio.to_thread(
                    index_single_markdown_file,
                    filepath, rel_path, client, store, dense_model, sparse_model,
                    markdown_splitter, parent_splitter, child_splitter
                )

            file_hash = await asyncio.to_thread(get_file_hash, filepath)
            state[rel_path] = {
                "hash": file_hash,
                "doc_ids": doc_ids,
                "point_ids": point_ids,
                "category": determine_category(rel_path),
                "updated_at": time.time()
            }
            await asyncio.to_thread(save_state, state)

        return {
            "status": "success",
            "relative_path": rel_path,
            "filename": fname,
            "chunks_count": len(doc_ids),
            "points_count": len(point_ids)
        }

    async def sync_all_documents(self) -> Dict[str, Any]:
        """Đồng bộ toàn bộ các thay đổi trong docs/ (thêm mới, chỉnh sửa, xóa) qua các subfolder."""
        if app_state.load_task and not app_state.load_task.done():
            await app_state.load_task

        client = app_state.qdrant_client
        store = app_state.store
        dense_model = app_state.dense_model
        sparse_model = app_state.sparse_model

        if not client or not store or not dense_model or not sparse_model:
            return {"status": "error", "reason": "Các thành phần AI chưa sẵn sàng."}

        async with indexing_lock:
            def _scan_and_hash():
                st = load_state()
                files = self.scan_all_files()
                curr = {}
                for item in files:
                    if is_indexable_doc(item["relative_path"]):
                        curr[item["relative_path"]] = {
                            "hash": get_file_hash(item["full_path"]),
                            "full_path": item["full_path"]
                        }
                return st, curr

            state, current_files = await asyncio.to_thread(_scan_and_hash)

            saved_keys = set(state.keys())

            added_keys = [k for k in current_files if k not in saved_keys]
            deleted_keys = [k for k in saved_keys if k not in current_files]
            modified_keys = [
                k for k in current_files
                if k in saved_keys and current_files[k]["hash"] != state[k].get("hash")
            ]

            # 1. Xóa file đã mất
            for k in deleted_keys:
                await asyncio.to_thread(delete_file_from_index, k, state[k], client, store)
                del state[k]

            # 2. Re-index file sửa đổi
            markdown_splitter, parent_splitter, child_splitter = init_splitters()
            for k in modified_keys:
                await asyncio.to_thread(delete_file_from_index, k, state[k], client, store)
                fpath = current_files[k]["full_path"]
                if k.lower().endswith((".xlsx", ".xls")):
                    doc_ids, point_ids = await asyncio.to_thread(
                        index_single_excel_file, fpath, k, client, store, dense_model, sparse_model
                    )
                else:
                    doc_ids, point_ids = await asyncio.to_thread(
                        index_single_markdown_file, fpath, k, client, store, dense_model, sparse_model,
                        markdown_splitter, parent_splitter, child_splitter
                    )
                state[k] = {
                    "hash": current_files[k]["hash"],
                    "doc_ids": doc_ids,
                    "point_ids": point_ids,
                    "category": determine_category(k),
                    "updated_at": time.time()
                }

            # 3. Index file thêm mới
            for k in added_keys:
                fpath = current_files[k]["full_path"]
                if k.lower().endswith((".xlsx", ".xls")):
                    doc_ids, point_ids = await asyncio.to_thread(
                        index_single_excel_file, fpath, k, client, store, dense_model, sparse_model
                    )
                else:
                    doc_ids, point_ids = await asyncio.to_thread(
                        index_single_markdown_file, fpath, k, client, store, dense_model, sparse_model,
                        markdown_splitter, parent_splitter, child_splitter
                    )
                state[k] = {
                    "hash": current_files[k]["hash"],
                    "doc_ids": doc_ids,
                    "point_ids": point_ids,
                    "category": determine_category(k),
                    "indexed_at": time.time()
                }

            await asyncio.to_thread(save_state, state)
            await asyncio.to_thread(refresh_domain_caches)

        return {
            "status": "success",
            "added": added_keys,
            "modified": modified_keys,
            "deleted": deleted_keys,
            "total_changes": len(added_keys) + len(modified_keys) + len(deleted_keys)
        }


doc_manager = DocumentManager()
